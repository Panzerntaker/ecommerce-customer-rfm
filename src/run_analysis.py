
import os
import sqlite3
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, "data")
SQL_DIR = os.path.join(BASE_DIR, "sql")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
DB_PATH = os.path.join(DATA_DIR, "ecommerce.db")


def get_db_connection() -> sqlite3.Connection:
    if not os.path.exists(DB_PATH):
        print("Database not found. Generating synthetic dataset...")
        from generate_data import generate_synthetic_data, save_data

        df = generate_synthetic_data()
        save_data(df)
    return sqlite3.connect(DB_PATH)


def run_sql_query(query_file: str, conn: sqlite3.Connection) -> pd.DataFrame:
    path = os.path.join(SQL_DIR, query_file)
    with open(path, "r", encoding="utf-8") as f:
        query = f.read()
    return pd.read_sql_query(query, conn)


def assign_rfm_segment(row: pd.Series) -> str:
    r = int(row["r_score"])
    f = int(row["f_score"])
    m = int(row["m_score"])

    if r >= 3 and f >= 3:
        if r == 4 and f >= 3 and m >= 3:
            return "Champions"
        return "Loyal"
    elif r >= 3 and f <= 2:
        if r == 4 and f == 1:
            return "Recent"
        return "Potential"
    elif r <= 2 and f >= 3:
        return "At Risk"
    else:  # r <= 2 and f <= 2
        return "Hibernating"


def generate_cohort_heatmap(df_cohort: pd.DataFrame, output_path: str):
    pivot_retention = df_cohort.pivot(
        index="cohort_month", columns="cohort_index", values="retention_rate"
    )

    plt.figure(figsize=(12, 7), dpi=300)
    sns.set_theme(style="white")

    cmap = sns.color_palette("Blues", as_cmap=True)
    ax = sns.heatmap(
        pivot_retention,
        annot=True,
        fmt=".1f",
        cmap=cmap,
        vmin=0,
        vmax=100,
        cbar_kws={"label": "Retention Rate (%)"},
        linewidths=0.5,
        linecolor="#e0e0e0",
    )

    ax.set_title(
        "Monthly Customer Cohort Retention Rate (%)",
        fontsize=14,
        fontweight="bold",
        pad=15,
    )
    ax.set_xlabel("Cohort Index (Months Since First Purchase)", fontsize=11, labelpad=10)
    ax.set_ylabel("Cohort Month", fontsize=11, labelpad=10)
    plt.xticks(rotation=0)
    plt.yticks(rotation=0)
    plt.tight_layout()
    plt.savefig(output_path)
    plt.close()
    print(f"Saved cohort retention plot: {output_path}")


def generate_rfm_plots(df_rfm: pd.DataFrame, output_path: str):
    segment_summary = (
        df_rfm.groupby("segment")
        .agg(
            customer_count=("customer_id", "count"),
            total_revenue=("monetary", "sum"),
            avg_recency=("recency_days", "mean"),
            avg_frequency=("frequency", "mean"),
        )
        .reset_index()
    )

    seg_order = ["Champions", "Loyal", "Potential", "Recent", "At Risk", "Hibernating"]
    segment_summary["segment"] = pd.Categorical(
        segment_summary["segment"], categories=seg_order, ordered=True
    )
    segment_summary = segment_summary.sort_values("segment")

    palette = {
        "Champions": "#1b7837",
        "Loyal": "#4575b4",
        "Potential": "#74add1",
        "Recent": "#abd9e9",
        "At Risk": "#f46d43",
        "Hibernating": "#d73027",
    }

    fig, axes = plt.subplots(1, 2, figsize=(15, 6), dpi=300)
    sns.set_theme(style="whitegrid")

    y = np.arange(len(segment_summary))
    height = 0.35

    ax1 = axes[0]
    bars1 = ax1.barh(
        y - height / 2,
        segment_summary["customer_count"],
        height,
        label="Customer Count",
        color="#3182bd",
        edgecolor="none",
    )
    ax1.set_xlabel("Customer Count", fontsize=11, fontweight="bold", color="#3182bd")
    ax1.set_yticks(y)
    ax1.set_yticklabels(segment_summary["segment"], fontsize=10, fontweight="bold")
    ax1.invert_yaxis()

    ax1_twin = ax1.twiny()
    bars2 = ax1_twin.barh(
        y + height / 2,
        segment_summary["total_revenue"] / 1000,
        height,
        label="Total Revenue ($k)",
        color="#31a354",
        edgecolor="none",
    )
    ax1_twin.set_xlabel("Total Revenue ($ in Thousands)", fontsize=11, fontweight="bold", color="#31a354")

    ax1.set_title("Customer Count & Revenue by Segment", fontsize=13, fontweight="bold", pad=15)

    for bar in bars1:
        w = bar.get_width()
        ax1.text(w + 1, bar.get_y() + bar.get_height() / 2, f"{int(w)}", va="center", fontsize=9, color="#1c4563")
    for bar in bars2:
        w = bar.get_width()
        ax1_twin.text(w + 0.5, bar.get_y() + bar.get_height() / 2, f"${w:.1f}k", va="center", fontsize=9, color="#1c6328")

    ax2 = axes[1]
    for seg in seg_order:
        subset = df_rfm[df_rfm["segment"] == seg]
        ax2.scatter(
            subset["recency_days"],
            subset["frequency"],
            s=subset["monetary"] / 25,
            color=palette[seg],
            alpha=0.65,
            edgecolors="w",
            linewidth=0.5,
            label=seg,
        )

    ax2.set_title("Recency vs. Frequency (Bubble Size = Spend)", fontsize=13, fontweight="bold", pad=15)
    ax2.set_xlabel("Recency (Days Since Last Order)", fontsize=11)
    ax2.set_ylabel("Frequency (Total Orders)", fontsize=11)
    ax2.legend(title="Segment", frameon=True, loc="upper right")

    plt.tight_layout()
    plt.savefig(output_path)
    plt.close()
    print(f"Saved RFM segmentation plot: {output_path}")


def export_excel_report(
    df_cohort: pd.DataFrame,
    df_rfm: pd.DataFrame,
    output_path: str,
):
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    title_font = Font(name="Calibri", size=15, bold=True, color="1F4E78")
    sub_font = Font(name="Calibri", size=11, italic=True, color="595959")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    kpi_val_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    kpi_lbl_font = Font(name="Calibri", size=10, color="595959")
    kpi_fill = PatternFill(start_color="EDF2F8", end_color="EDF2F8", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws_exec = wb.create_sheet(title="Executive Summary")
    ws_exec.views.sheetView[0].showGridLines = True

    ws_exec["A1"] = "E-Commerce Customer Retention & RFM Executive Report"
    ws_exec["A1"].font = title_font
    ws_exec["A2"] = "Comprehensive portfolio analytics: Retention dynamics, RFM customer tiers, and strategic actions."
    ws_exec["A2"].font = sub_font

    total_customers = df_rfm["customer_id"].nunique()
    total_revenue = df_rfm["monetary"].sum()
    total_orders = df_rfm["frequency"].sum()
    aov = total_revenue / total_orders
    clv = total_revenue / total_customers

    m1_cohorts = df_cohort[df_cohort["cohort_index"] == 1]
    avg_m1_retention = m1_cohorts["retention_rate"].mean() if len(m1_cohorts) > 0 else 0.0

    kpis = [
        ("Total Customers", f"{total_customers:,}"),
        ("Total Revenue", f"${total_revenue:,.2f}"),
        ("Total Orders", f"{total_orders:,}"),
        ("Average Order Value (AOV)", f"${aov:,.2f}"),
        ("Customer Lifetime Value (CLV)", f"${clv:,.2f}"),
        ("Average Month-1 Retention", f"{avg_m1_retention:.1f}%"),
    ]

    ws_exec["A4"] = "KEY PERFORMANCE INDICATORS"
    ws_exec["A4"].font = bold_font

    for idx, (label, val) in enumerate(kpis):
        col_start = (idx % 3) * 3 + 1
        row_start = 5 if idx < 3 else 8

        c1 = ws_exec.cell(row=row_start, column=col_start, value=label)
        c1.font = kpi_lbl_font
        c1.fill = kpi_fill
        c1.alignment = Alignment(horizontal="center", vertical="center")

        c2 = ws_exec.cell(row=row_start + 1, column=col_start, value=val)
        c2.font = kpi_val_font
        c2.fill = kpi_fill
        c2.alignment = Alignment(horizontal="center", vertical="center")

        ws_exec.merge_cells(start_row=row_start, start_column=col_start, end_row=row_start, end_column=col_start + 2)
        ws_exec.merge_cells(start_row=row_start + 1, start_column=col_start, end_row=row_start + 1, end_column=col_start + 2)

    ws_exec["A11"] = "STRATEGIC RECOMMENDATIONS BY SEGMENT"
    ws_exec["A11"].font = bold_font

    strategy_headers = ["Segment", "Base Share", "Revenue Share", "Key Behavioral Profile", "Actionable Recommendation"]
    for col_idx, h in enumerate(strategy_headers, start=1):
        cell = ws_exec.cell(row=12, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    seg_summary = (
        df_rfm.groupby("segment")
        .agg(
            count=("customer_id", "count"),
            rev=("monetary", "sum"),
        )
        .reset_index()
    )
    seg_summary["base_pct"] = seg_summary["count"] / total_customers * 100
    seg_summary["rev_pct"] = seg_summary["rev"] / total_revenue * 100

    strat_map = {
        "Champions": ("Bought recently, buy often, and spend the most.", "VIP loyalty tier, early access to releases, personal concierge."),
        "Loyal": ("Regular shoppers with solid spend and consistent frequency.", "Upsell higher margin categories, subscription loyalty programs."),
        "Potential": ("Recent purchasers with moderate frequency and spend.", "Targeted cross-sell sequences and gamified loyalty milestones."),
        "Recent": ("First-time or recent buyers who made single purchases.", "Automated onboarding email drip with welcome-back discount."),
        "At Risk": ("High past spenders and frequent buyers who haven't ordered recently.", "Win-back campaigns, personalized surveys, aggressive incentives."),
        "Hibernating": ("Low frequency, low spend, last order was long ago.", "Low-cost re-engagement via push/email or automated sunsetting."),
    }

    for idx, row in seg_summary.iterrows():
        seg_name = row["segment"]
        profile, action = strat_map.get(seg_name, ("", ""))
        r_idx = 13 + idx
        ws_exec.cell(row=r_idx, column=1, value=seg_name).font = bold_font
        ws_exec.cell(row=r_idx, column=2, value=f"{row['base_pct']:.1f}%").alignment = Alignment(horizontal="right")
        ws_exec.cell(row=r_idx, column=3, value=f"{row['rev_pct']:.1f}%").alignment = Alignment(horizontal="right")
        ws_exec.cell(row=r_idx, column=4, value=profile)
        ws_exec.cell(row=r_idx, column=5, value=action)
        for col_idx in range(1, 6):
            ws_exec.cell(row=r_idx, column=col_idx).border = thin_border

    ws_rfm_sum = wb.create_sheet(title="RFM Segment Summary")
    ws_rfm_sum.views.sheetView[0].showGridLines = True

    summary_df = (
        df_rfm.groupby("segment")
        .agg(
            customer_count=("customer_id", "count"),
            total_spend=("monetary", "sum"),
            avg_spend=("monetary", "mean"),
            avg_orders=("frequency", "mean"),
            avg_recency_days=("recency_days", "mean"),
        )
        .reset_index()
    )
    summary_df["pct_customers"] = summary_df["customer_count"] / total_customers
    summary_df["pct_revenue"] = summary_df["total_spend"] / total_revenue

    cols_order = [
        ("Segment", "segment", "@"),
        ("Customer Count", "customer_count", "#,##0"),
        ("% of Base", "pct_customers", "0.0%"),
        ("Total Spend", "total_spend", "$#,##0.00"),
        ("% of Revenue", "pct_revenue", "0.0%"),
        ("Avg Customer Spend", "avg_spend", "$#,##0.00"),
        ("Avg Order Count", "avg_orders", "0.0"),
        ("Avg Recency (Days)", "avg_recency_days", "0.0"),
    ]

    for c_idx, (col_name, _, _) in enumerate(cols_order, start=1):
        cell = ws_rfm_sum.cell(row=1, column=c_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, row in summary_df.iterrows():
        current_r = r_idx + 2
        for c_idx, (_, field, num_fmt) in enumerate(cols_order, start=1):
            val = row[field]
            cell = ws_rfm_sum.cell(row=current_r, column=c_idx, value=val)
            cell.font = regular_font
            cell.border = thin_border
            cell.number_format = num_fmt
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")

    tot_row = len(summary_df) + 2
    ws_rfm_sum.cell(row=tot_row, column=1, value="Total").font = bold_font
    ws_rfm_sum.cell(row=tot_row, column=2, value=f"=SUM(B2:B{tot_row-1})").number_format = "#,##0"
    ws_rfm_sum.cell(row=tot_row, column=3, value=f"=SUM(C2:C{tot_row-1})").number_format = "0.0%"
    ws_rfm_sum.cell(row=tot_row, column=4, value=f"=SUM(D2:D{tot_row-1})").number_format = "$#,##0.00"
    ws_rfm_sum.cell(row=tot_row, column=5, value=f"=SUM(E2:E{tot_row-1})").number_format = "0.0%"
    ws_rfm_sum.cell(row=tot_row, column=6, value=f"=AVERAGE(F2:F{tot_row-1})").number_format = "$#,##0.00"
    ws_rfm_sum.cell(row=tot_row, column=7, value=f"=AVERAGE(G2:G{tot_row-1})").number_format = "0.0"
    ws_rfm_sum.cell(row=tot_row, column=8, value=f"=AVERAGE(H2:H{tot_row-1})").number_format = "0.0"
    for c_idx in range(1, 9):
        c = ws_rfm_sum.cell(row=tot_row, column=c_idx)
        c.font = bold_font
        c.border = thin_border

    ws_rfm_det = wb.create_sheet(title="Customer RFM Details")
    ws_rfm_det.views.sheetView[0].showGridLines = True

    det_cols = [
        ("Customer ID", "customer_id", "@"),
        ("Recency (Days)", "recency_days", "#,##0"),
        ("Frequency (Orders)", "frequency", "#,##0"),
        ("Monetary Spend", "monetary", "$#,##0.00"),
        ("R Score", "r_score", "0"),
        ("F Score", "f_score", "0"),
        ("M Score", "m_score", "0"),
        ("RFM Cell", "rfm_cell", "@"),
        ("Segment", "segment", "@"),
    ]

    for c_idx, (col_name, _, _) in enumerate(det_cols, start=1):
        cell = ws_rfm_det.cell(row=1, column=c_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, row in df_rfm.iterrows():
        current_r = r_idx + 2
        for c_idx, (_, field, num_fmt) in enumerate(det_cols, start=1):
            val = row[field]
            cell = ws_rfm_det.cell(row=current_r, column=c_idx, value=val)
            cell.font = regular_font
            cell.number_format = num_fmt
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right")

    ws_cohort = wb.create_sheet(title="Cohort Retention Matrix")
    ws_cohort.views.sheetView[0].showGridLines = True

    pivot_ret = df_cohort.pivot(index="cohort_month", columns="cohort_index", values="retention_rate")
    cohort_sizes = df_cohort.groupby("cohort_month")["cohort_size"].first()

    ws_cohort["A1"] = "Cohort Month"
    ws_cohort["A1"].font = header_font
    ws_cohort["A1"].fill = header_fill

    ws_cohort["B1"] = "Cohort Size"
    ws_cohort["B1"].font = header_font
    ws_cohort["B1"].fill = header_fill

    max_idx = int(df_cohort["cohort_index"].max())
    for idx_num in range(max_idx + 1):
        cell = ws_cohort.cell(row=1, column=3 + idx_num, value=f"Month {idx_num}")
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, (month, row_vals) in enumerate(pivot_ret.iterrows(), start=2):
        ws_cohort.cell(row=r_idx, column=1, value=month).font = bold_font
        ws_cohort.cell(row=r_idx, column=2, value=int(cohort_sizes[month])).number_format = "#,##0"

        for c_idx, val in enumerate(row_vals, start=3):
            cell = ws_cohort.cell(row=r_idx, column=c_idx)
            if pd.notnull(val):
                cell.value = val / 100.0
                cell.number_format = "0.0%"
                cell.alignment = Alignment(horizontal="right")
                if val >= 50:
                    cell.fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
                elif val >= 30:
                    cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
            cell.border = thin_border

    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_path)
    print(f"Saved executive Excel report: {output_path}")


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    conn = get_db_connection()

    print("Running Cohort Retention Analysis (SQL)...")
    df_cohort = run_sql_query("01_cohort_retention.sql", conn)

    print("Running RFM Scoring Analysis (SQL)...")
    df_rfm = run_sql_query("02_rfm_scoring.sql", conn)

    df_rfm["segment"] = df_rfm.apply(assign_rfm_segment, axis=1)

    cohort_img = os.path.join(OUTPUT_DIR, "cohort_retention.png")
    rfm_img = os.path.join(OUTPUT_DIR, "rfm_segments.png")
    excel_report = os.path.join(OUTPUT_DIR, "ecommerce_executive_report.xlsx")

    generate_cohort_heatmap(df_cohort, cohort_img)
    generate_rfm_plots(df_rfm, rfm_img)
    export_excel_report(df_cohort, df_rfm, excel_report)

    conn.close()
    print("Analysis pipeline execution completed successfully.")


if __name__ == "__main__":
    main()
